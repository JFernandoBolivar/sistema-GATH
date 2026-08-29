"""
Tests unitarios del motor genérico de reportes Excel (Fase 2.4 del plan de
reparación, `PLAN_REPARACION_GATH.md`). Puro: openpyxl no depende de
Django ni de la BD.
"""
from io import BytesIO

from openpyxl import load_workbook

from core.infrastructure.pdf.report_excel import generar_reporte_excel


def test_generar_reporte_excel_con_filas():
    excel_bytes = generar_reporte_excel(
        titulo="Reporte de Personal Activo",
        columnas=['cedula', 'nombres', 'apellidos'],
        filas=[
            {'cedula': '12345678', 'nombres': 'Juan', 'apellidos': 'Pérez'},
            {'cedula': '87654321', 'nombres': 'Ana', 'apellidos': 'Gómez'},
        ],
    )
    wb = load_workbook(BytesIO(excel_bytes))
    ws = wb.active
    assert ws['A1'].value == 'CEDULA'
    assert ws['A2'].value == '12345678'
    assert ws['B2'].value == 'Juan'
    assert ws.max_row == 3


def test_generar_reporte_excel_sin_filas_solo_encabezado():
    excel_bytes = generar_reporte_excel(titulo="Vacío", columnas=['id', 'nombre'], filas=[])
    wb = load_workbook(BytesIO(excel_bytes))
    ws = wb.active
    assert ws.max_row == 1
    assert ws['A1'].value == 'ID'


def test_generar_reporte_excel_valor_faltante_se_deja_vacio():
    excel_bytes = generar_reporte_excel(
        titulo="Con nulos", columnas=['a', 'b'], filas=[{'a': 1}],
    )
    wb = load_workbook(BytesIO(excel_bytes))
    ws = wb.active
    assert ws['A2'].value == 1
    assert ws['B2'].value is None
