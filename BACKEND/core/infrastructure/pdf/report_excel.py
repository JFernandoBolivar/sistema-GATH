"""
Motor genérico de reportes Excel (Fase 2.4 del plan de reparación,
`PLAN_REPARACION_GATH.md`). Vive junto a `report_pdf.py` porque comparte el
mismo insumo genérico (columnas/filas de un `ResultadoReporte`) — no hay un
motor Excel legado que portar (el único que existía era
`AUTOGESTION/services/excel_service.py`, específico del censo, ver Fase 2.6).
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def generar_reporte_excel(
    titulo: str, columnas: list[str], filas: list[dict],
    filtros: list[tuple[str, str]] | None = None,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = titulo[:31] or 'Reporte'

    encabezado_fill = PatternFill(start_color='1A237E', end_color='1A237E', fill_type='solid')
    encabezado_font = Font(color='FFFFFF', bold=True)

    fila_encabezado = 1
    if filtros:
        for etiqueta, valor in filtros:
            ws.cell(row=fila_encabezado, column=1, value=f'{etiqueta}:').font = Font(bold=True)
            ws.cell(row=fila_encabezado, column=2, value=str(valor))
            fila_encabezado += 1
        fila_encabezado += 1  # fila en blanco antes de la tabla

    for col_idx, columna in enumerate(columnas, start=1):
        celda = ws.cell(row=fila_encabezado, column=col_idx, value=columna.replace('_', ' ').upper())
        celda.fill = encabezado_fill
        celda.font = encabezado_font

    for fila_idx, fila in enumerate(filas, start=fila_encabezado + 1):
        for col_idx, columna in enumerate(columnas, start=1):
            ws.cell(row=fila_idx, column=col_idx, value=fila.get(columna))

    for col_idx in range(1, len(columnas) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 22

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
